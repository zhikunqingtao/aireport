# -*- coding: utf-8 -*-
from openpyxl import load_workbook
OUT="/Users/guoqingtao/Desktop/AI办公评测/豆包办公/宇树科技经营财务分析底稿_截至20260830.xlsx"
wb=load_workbook(OUT)  # formulas
print("SHEETS:",wb.sheetnames)
# 1 图表数量
wc=wb["07_趋势图表"]; print("CHARTS on 07:",len(wc._charts))
# 2 网格线
for s in wb.sheetnames:
    assert wb[s].sheet_view.showGridLines==False, s
print("gridlines off: OK")
# 3 独立复算（与公式应得结果对照）
rev={2023:15913.44,2024:39277.07,2025:169926.93}
ni={2023:-1114.51,2024:9547.47,2025:27821.05}
nd={2023:-1801.91,2024:7847.65,2025:59075.28}
cagr=(rev[2025]/rev[2023])**0.5-1
print("CAGR=%.4f (disclose 226.78pct)"%cagr)
for y in (2024,2025):
    print(y,"营收同比=%.2f%%  扣非净利率=%.2f%%  销售净利率=%.2f%%"%(
        (rev[y]/rev[y-1]-1)*100, nd[y]/rev[y]*100, ni[y]/rev[y]*100))
# 勾稽：分产品/分地区合计
prod23=296.71+11938.09+2692.34+826.51
prod24=10689.76+23054.37+4453.82+529.33
print("分产品合计 2023=%.2f(应15753.65) 2024=%.2f(应38727.28)"%(prod23,prod24))
reg23=6989.41+8764.24; reg24=17116.54+21610.74; reg25=94445.56+73165.53
print("分地区合计 2023=%.2f 2024=%.2f 2025=%.2f(应167611.09)"%(reg23,reg24,reg25))
print("2025非经常性=%.2f (应-31254.23)"%(ni[2025]-nd[2025]))
# 2026H1 同比
h1_26=115224.56; h1_25=77571.4
print("2026H1营收同比=%.2f%% (披露48.54)"%((h1_26/h1_25-1)*100))
print("2026H1扣非同比=%.2f%% (披露-19.34)"%((24393.03/30243.35-1)*100))
# 单季
q=[("25Q2",77571.4-25095.9),("25Q3",116749.01-77571.4),("25Q4",169926.93-116749.01),("26Q2",115224.56-42284.05)]
print("单季营收:",[(a,round(b,2)) for a,b in q])
# 4 扫描公式错误标记 & 统计公式数
bad=[]; nf=0
for s in wb.sheetnames:
    for row in wb[s].iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str):
                if v.startswith("="): nf+=1
                if any(e in str(v) for e in ["#REF!","#DIV/0!","#NAME?","#VALUE!","#N/A"]): bad.append((s,c.coordinate,v))
print("formula cells:",nf,"| error tokens:",bad)
# 5 抽查关键公式文本
w2=wb["02_原始数据_年度"]; print("02!E28 研发占比公式:",w2["E28"].value)
w5=wb["05_核心计算"]
for rr in range(4,19):
    if w5[f"B{rr}"].value: print("05 row",rr,w5[f"B{rr}"].value,"| E:",w5[f"E{rr}"].value)

# -*- coding: utf-8 -*-
import re
from openpyxl import load_workbook
OUT="宇树科技经营财务分析底稿_截至20260830.xlsx"
wb=load_workbook(OUT)
names=set(wb.sheetnames)
ref_re=re.compile(r"'([^']+)'!|([0-9A-Za-z_\u4e00-\u9fff]+)!")
cell_re=re.compile(r"(?:'([^']+)'|([0-9A-Za-z_\u4e00-\u9fff]+))!\$?([A-Z]{1,3})\$?(\d+)")
problems=[]; total=0
for s in wb.sheetnames:
    ws=wb[s]
    for row in ws.iter_rows():
        for c in row:
            v=c.value
            if not (isinstance(v,str) and v.startswith("=")): continue
            total+=1
            # 括号平衡
            if v.count("(")!=v.count(")"): problems.append((s,c.coordinate,"paren",v))
            # 跨表引用目标存在且非空
            for m in cell_re.finditer(v):
                sh=m.group(1) or m.group(2); col=m.group(3); rr=int(m.group(4))
                if sh not in names: problems.append((s,c.coordinate,"no-sheet:"+sh,v)); continue
                tgt=wb[sh][f"{col}{rr}"]
                if tgt.value is None:
                    problems.append((s,c.coordinate,f"empty-target {sh}!{col}{rr}",v))
print("total formulas:",total)
print("problems:",len(problems))
for p in problems[:40]: print(p)
